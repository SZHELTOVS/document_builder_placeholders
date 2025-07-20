import os
from docxtpl import DocxTemplate
from django.utils.crypto import get_random_string
from docxtpl import DocxTemplate
import docx
import re

def extract_placeholders(docx_path):
    doc = DocxTemplate(docx_path)
    context = list(doc.get_undeclared_template_variables())

    docx_doc = docx.Document(docx_path)
    tables = []

    for table in docx_doc.tables:
        text = '\n'.join(cell.text for row in table.rows for cell in row.cells)
        match = re.search(r"{%\s*for\s+row\s+in\s+(\w+)\s*%}", text)
        if match:
            table_name = match.group(1)

            # Собираем колонки из текста всех ячеек (сохраняем порядок)
            cols = []
            for row in table.rows:
                for cell in row.cells:
                    cols += re.findall(r"{{\s*row\.([\w_]+)\s*}}", cell.text)
            columns = list(dict.fromkeys(cols))

            if columns:
                tables.append({
                    "type": "table",
                    "name": table_name,
                    "columns": [{"name": c, "label": c.replace('_', ' ').capitalize()} for c in columns]
                })

    table_names = [tbl['name'] for tbl in tables]
    print(table_names)
    print(len(table_names), len(set(table_names)))  # должны совпадать
    
    simple_fields = [f for f in context if not any(tbl['name'] == f for tbl in tables)]

    return simple_fields + tables

def replace_placeholders(path, values):
    doc = DocxTemplate(path)
    doc.render(values)

    output_filename = f"filled_{get_random_string(10)}.docx"
    output_path = os.path.join("media/output", output_filename)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    doc.save(output_path)
    print(f"==== SAVED TO: {output_path} EXISTS: {os.path.exists(output_path)} SIZE: {os.path.getsize(output_path)}")
    return output_path

